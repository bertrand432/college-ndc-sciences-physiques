from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(__file__).parent
def base(title):
 d=Document(); s=d.sections[0]; s.page_height=Cm(29.7); s.page_width=Cm(21)
 s.top_margin=s.bottom_margin=Cm(1.7); s.left_margin=s.right_margin=Cm(1.8)
 for n in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[n]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
 st=d.styles['Normal']; st.font.size=Pt(12); st.paragraph_format.space_after=Pt(7)
 st.paragraph_format.line_spacing=1.12
 d.styles['Title'].font.size=Pt(22); d.styles['Heading 1'].font.size=Pt(15)
 d.add_paragraph(title,'Title')
 return d
def p(d,t): return d.add_paragraph(t)
def h(d,t): d.add_paragraph(t,'Heading 1')
def lines(d,n=1):
 for _ in range(n): p(d,'................................................................................................')
def table(d,rows):
 t=d.add_table(rows=0,cols=len(rows[0])); t.autofit=False
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for c,txt in zip(cells,row):
   c.text=txt
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   if i==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EDEDED'); pr.append(sh)
   for para in c.paragraphs:
    para.paragraph_format.space_before=Pt(5); para.paragraph_format.space_after=Pt(5)
    for r in para.runs: r.bold=i==0
 return t

d=base('Rappels en électricité')
p(d,'Sciences physiques • 4e SEGPA • Chapitre 1 Mesures électriques')
p(d,'Nom : ...........................................  Date : ........................')
p(d,'Je reconnais le matériel et je comprends un circuit simple.\nLis une consigne à la fois. Tu peux utiliser ton cahier.')
h(d,'1 Le rôle du matériel')
p(d,'Complète avec : pile • lampe • interrupteur • fils électriques.')
p(d,'a. La ................................ fournit de l’énergie électrique.')
p(d,'b. La ................................ produit de la lumière.')
p(d,'c. L’................................ permet d’ouvrir ou de fermer le circuit.')
p(d,'d. Les ................................ relient les éléments du circuit.')
h(d,'2 La lampe peut elle briller')
p(d,'La pile et la lampe fonctionnent. Entoure la bonne réponse.')
table(d,[['Situation','La lampe brille'],['La boucle est complète et l’interrupteur est fermé.','OUI / NON'],['L’interrupteur est ouvert.','OUI / NON'],['Un fil est débranché.','OUI / NON']])
p(d,'Complète avec : ouvert • fermé.')
p(d,'Le courant peut circuler dans un circuit ................................ .')
p(d,'Si la boucle est coupée, le circuit est ................................ .')
h(d,'3 Conducteur ou isolant')
p(d,'Un conducteur laisse passer le courant. Un isolant ne le laisse pas passer dans les conditions habituelles de la classe.')
p(d,'Écris C pour conducteur ou I pour isolant.')
p(d,'Cuivre : ........     Plastique : ........     Fer : ........     Verre : ........')
p(d,'Dans un fil électrique, le métal est entouré de plastique.\nÀ quoi sert ce plastique ?')
lines(d)
d.add_page_break()
d.add_paragraph('Les premières mesures électriques','Title')
p(d,'Je distingue la tension et l’intensité et je lis une mesure.')
p(d,'Aide : la tension se mesure avec un voltmètre, en volts (V).\nL’intensité se mesure avec un ampèremètre, en ampères (A).')
h(d,'4 Choisir une mesure')
p(d,'Complète le tableau avec les mots de l’aide.')
table(d,[['Grandeur mesurée','Appareil','Unité et symbole'],['Tension','........................','........................'],['Intensité','........................','........................']])
p(d,'Entoure la bonne réponse.')
p(d,'a. Sur une pile, « 4,5 V » indique : la tension / l’intensité.')
p(d,'b. « 0,20 A » est une mesure de : tension / intensité.')
h(d,'5 Lire les appareils')
p(d,'Écris chaque mesure avec son unité. Exemple : 3,0 V.')
table(d,[['Appareil','Affichage','Mesure avec unité'],['Voltmètre','1,5','........................'],['Voltmètre','4,5','........................'],['Ampèremètre','0,30','........................']])
p(d,'Deux piles ont des tensions de 1,5 V et de 4,5 V.\nEntoure la tension la plus grande : 1,5 V / 4,5 V.')
h(d,'6 Décrire un branchement')
p(d,'Aide : un ampèremètre se place en série, dans la boucle.\nUn voltmètre se place en dérivation, aux deux bornes de la lampe.')
p(d,'Complète avec : ampèremètre • voltmètre.')
p(d,'a. Pour mesurer la tension aux bornes de la lampe,\nj’utilise un ........................................................ .')
p(d,'b. Pour mesurer l’intensité du courant dans la lampe,\nj’utilise un ........................................................ .')
p(d,'Entoure : le voltmètre se branche en série / en dérivation.')
p(d,'Au laboratoire : réalise les montages avec les piles fournies et fais vérifier les branchements par le professeur avant de fermer le circuit.')
d.save(OUT/'Exercices_electricite_4e_SEGPA.docx')

c=base('Corrigé des exercices en électricité')
p(c,'Sciences physiques • 4e SEGPA • Document enseignant')
p(c,'Objectif : consolider les bases du circuit électrique et découvrir les mesures de tension et d’intensité.')
h(c,'1 Le rôle du matériel')
p(c,'a. pile ; b. lampe ; c. interrupteur ; d. fils électriques.')
h(c,'2 La lampe peut elle briller')
p(c,'Dans l’ordre du tableau : OUI ; NON ; NON.\nMots à compléter : fermé ; ouvert.')
h(c,'3 Conducteur ou isolant')
p(c,'Cuivre : C ; plastique : I ; fer : C ; verre : I.')
p(c,'Réponse attendue : le plastique isole le métal. Il évite le contact direct avec le conducteur et le contact entre les conducteurs de différents fils.')
h(c,'4 Choisir une mesure')
p(c,'Tension : voltmètre ; volt (V).\nIntensité : ampèremètre ; ampère (A).\na. la tension ; b. intensité.')
h(c,'5 Lire les appareils')
p(c,'Dans l’ordre du tableau : 1,5 V ; 4,5 V ; 0,30 A.\nLa tension la plus grande est 4,5 V.')
h(c,'6 Décrire un branchement')
p(c,'a. voltmètre ; b. ampèremètre.\nLe voltmètre se branche en dérivation.')
h(c,'Pistes pour accompagner la séance')
p(c,'Prévoir environ 30 à 40 minutes selon l’aide apportée. Lire les consignes à voix haute si nécessaire. Montrer le matériel réel avant l’exercice 1. Les exercices 4 à 6 peuvent être réalisés ensemble si ces notions sont nouvelles.')
p(c,'Accepter les réponses orales avant le passage à l’écrit. Pour prolonger, faire montrer les deux bornes de la lampe sur un montage à pile, puis placer le voltmètre avec l’aide du professeur.')
c.save(OUT/'Corrige_electricite_4e_SEGPA.docx')
print('Deux documents Word créés.')
