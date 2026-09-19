from pathlib import Path
import json, shutil
import pypdfium2 as pdf
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path.cwd().parent
OUT=Path.cwd()/'output'
STAGE=OUT/'site-chimie'
rows=[
('Quel gaz permet au carbone de brûler ?', ['Le dioxyde de carbone.','Le dioxygène.','La vapeur d’eau.'],1,'Le dioxygène permet la combustion du carbone.'),
('Quel gaz se forme quand le carbone brûle ?', ['Le dioxyde de carbone.','Le dioxygène.','La vapeur d’eau.'],0,'La combustion transforme le carbone en dioxyde de carbone.'),
('Quel liquide permet de détecter le dioxyde de carbone ?', ['L’eau du robinet.','L’eau salée.','L’eau de chaux.'],2,'Dans le cours, on utilise l’eau de chaux pour détecter ce gaz.'),
('Que se passe-t-il quand l’eau de chaux rencontre du dioxyde de carbone ?', ['Elle se trouble.','Elle devient un gaz.','Elle devient un solide dur.'],0,'L’eau de chaux devient trouble en présence de dioxyde de carbone.'),
('Lors d’une transformation chimique, que deviennent les réactifs ?', ['Ils restent tous identiques.','Ils se transforment en nouvelles substances.','Ils deviennent toujours de l’eau.'],1,'De nouvelles substances apparaissent : ce sont les produits.'),
('Comment appelle-t-on les substances qui se transforment pendant la réaction ?', ['Les produits.','Les états.','Les réactifs.'],2,'Les réactifs sont les substances qui participent à la réaction et se transforment.'),
('Comment appelle-t-on les nouvelles substances formées ?', ['Les produits.','Les réactifs.','Les combustions.'],0,'Les produits sont les nouvelles substances obtenues après la réaction.'),
('Quel exemple de transformation chimique est donné dans le cours ?', ['Observer une dune de sable.','Faire brûler du carbone.','Représenter une particule.'],1,'Quand le carbone brûle, une nouvelle substance se forme : le dioxyde de carbone.'),
('Dans un bilan de réaction, que place-t-on avant la flèche ?', ['Les noms des produits.','Les noms des états de la matière.','Les noms des réactifs.'],2,'On écrit les réactifs avant la flèche et les produits après.'),
('Quel bilan correspond à la combustion du carbone ?', ['Carbone + dioxyde de carbone → dioxygène.','Carbone + dioxygène → dioxyde de carbone.','Dioxygène + eau de chaux → carbone.'],1,'Le carbone réagit avec le dioxygène pour former du dioxyde de carbone.'),
('Dans l’exemple du cours, de quoi une dune est-elle formée ?', ['De grains de sable.','De morceaux de papier.','De gouttes d’eau.'],0,'La dune est composée de nombreux grains de sable.'),
('Que permet de comprendre l’exemple de la dune ?', ['Toute la matière est du sable.','Toute la matière est liquide.','La matière est formée de particules.'],2,'Comme la dune est faite de grains, la matière est faite de particules.'),
('Quels noms donne-t-on aux particules dans le cours ?', ['Solides ou liquides.','Atomes ou molécules.','Dunes ou rochers.'],1,'Le cours appelle ces particules des atomes ou des molécules.'),
('Quels sont les trois états de la matière cités dans le cours ?', ['Solide, liquide et gazeux.','Solide, liquide et compressible.','Solide, gazeux et horizontal.'],0,'La matière peut être à l’état solide, liquide ou gazeux.'),
('Quel état de la matière possède une forme propre ?', ['L’état gazeux.','L’état liquide.','L’état solide.'],2,'Un solide possède sa propre forme.'),
('Quel état occupe tout le volume du récipient qui le contient ?', ['L’état liquide.','L’état gazeux.','L’état solide.'],1,'Un gaz occupe tout le volume disponible dans le récipient.'),
('Au repos, comment est la surface libre d’un liquide ?', ['Plane et horizontale.','Toujours inclinée.','Toujours bombée.'],0,'Quand le liquide est au repos, sa surface libre est plate et horizontale.'),
('D’après le tableau, que fait un liquide dans un récipient ?', ['Il garde toujours sa propre forme.','Il remplit toujours tout le volume du récipient.','Il prend la forme du récipient.'],2,'Un liquide n’a pas de forme propre : il prend la forme du récipient qui le contient.'),
('Quel état est indiqué comme « impalpable » dans le tableau ?', ['L’état solide.','L’état gazeux.','L’état liquide.'],1,'Le tableau indique « oui » pour les gaz : on ne peut pas les saisir dans la main.'),
('Quel état de la matière est compressible ?', ['L’état gazeux.','L’état solide.','L’état liquide.'],0,'Un gaz est compressible : on peut réduire le volume qu’il occupe.')]

for folder in ['js','css','qcm/3e','images/qcm-3d-chimie']:
    (STAGE/folder).mkdir(parents=True,exist_ok=True)
images=STAGE/'images/qcm-3d-chimie'
source=pdf.PdfDocument(str(ROOT/'documents/3d/cours1_2026.pdf'))
# Actual course crops at high resolution; hide only the answer on the reaction diagram.
im=source[3].render(scale=3).to_pil().convert('RGB')
im=im.crop((int(40*3),int(100*3),int(740*3),int(212*3)))
ImageDraw.Draw(im).rectangle((0,0,300*3,112*3),fill=(153,153,153))
im.save(images/'bilan-a-completer.png')
table=source[8].render(scale=3).to_pil().convert('RGB')
for name,y0,y1 in [('volume',165,205),('forme-liquide',265,306),('impalpable',306,327)]:
    head=table.crop((44*3,100*3,704*3,143*3))
    line=table.crop((44*3,y0*3,704*3,y1*3))
    crop=Image.new('RGB',(head.width,head.height+line.height),'white')
    crop.paste(head,(0,0)); crop.paste(line,(0,head.height))
    crop.save(images/(name+'.png'))
visuals={8:('bilan-a-completer.png','Bilan : une zone à compléter avant la flèche, puis les noms des produits.'),15:('volume.png','Tableau du cours : occuper tout le volume du récipient. Solide : non ; liquide : non ; gaz : oui.'),17:('forme-liquide.png','Tableau du cours : prendre la forme du récipient. Solide : non ; liquide : oui ; gaz : non.'),18:('impalpable.png','Tableau du cours : impalpable. Solide : non ; liquide : non ; gaz : oui.')}
questions=[]
for i,(question,answers,correct,explanation) in enumerate(rows):
    item=dict(question=question,answers=answers,correct=correct,explanation=explanation,category='Combustions et transformations chimiques' if i<10 else 'Le modèle particulaire' if i<13 else 'Les états de la matière')
    if i in visuals:
        item.update(image='../../images/qcm-3d-chimie/'+visuals[i][0],imageAlt=visuals[i][1])
    questions.append(item)
engine=(ROOT/'js/quiz-3d-securite.js').read_text(encoding='utf-8').split('// Each round')[1]
engine=engine.replace('Tu connais les dangers et les gestes de protection.','Tu connais les transformations chimiques et la matière.').replace('Revois les pictogrammes et les gestes de protection','Revois les transformations chimiques et les états de la matière')
engine=engine.replace('imageEl.hidden = !q.image;', 'imageEl.hidden = !q.image;\n  document.getElementById("courseVisual").hidden = !q.image;')
(STAGE/'js/quiz-3d-chimie.js').write_text('// Questions validées à partir de cours1_2026.pdf.\nconst questions = '+json.dumps(questions,ensure_ascii=False,indent=2)+';\n\n// Each round'+engine,encoding='utf-8')
html=(ROOT/'qcm/3e/securite-3d.html').read_text(encoding='utf-8').replace('20 questions pour les 3D SEGPA : pictogrammes de danger, risques et moyens de protection.','20 questions pour les 3D SEGPA : transformations chimiques, particules et états de la matière.').replace('Dangers et protections','Transformations chimiques et matière').replace('20 questions pour reconnaître les pictogrammes et apprendre à se protéger.','20 questions sur les combustions, les particules et les états de la matière.').replace('quiz-3d-securite.js','quiz-3d-chimie.js').replace('</head>','  <link rel="stylesheet" href="../../css/quiz-3d-chimie.css">\n</head>')
html=html.replace('<img id="pictogramme" class="quiz-pictogramme" alt="" hidden>', '<div id="courseVisual" hidden><p class="image-help">Fais défiler l’image horizontalement pour lire le tableau.</p><div class="image-scroll" tabindex="0" role="region" aria-label="Extrait du cours"><img id="pictogramme" class="quiz-pictogramme" alt="" hidden></div></div>')
(STAGE/'qcm/3e/chimie-3d.html').write_text(html,encoding='utf-8')
(STAGE/'css/quiz-3d-chimie.css').write_text('.quiz-pictogramme { width: 700px; max-width: 100%; height: auto; margin-bottom: 0; }\n.image-scroll { overflow-x: auto; margin-bottom: 24px; }\n.image-help { display: none; }\n@media(max-width: 700px) { .quiz-pictogramme { min-width: 640px; } .image-help { display: block; font-size: .9rem; } }\n',encoding='utf-8')
index=(ROOT/'qcm/3e/index.html').read_text(encoding='utf-8')
card='''
      <a class="chapter-card" href="chimie-3d.html">
        <span class="section-kicker">Disponible • 3D SEGPA</span>
        <h3>Transformations chimiques et matière</h3>
        <p>20 questions sur les combustions, les particules et les états de la matière.</p>
        <span class="card-link">Commencer →</span>
      </a>'''
assert 'chimie-3d.html' not in index
(STAGE/'qcm/3e/index.html').write_text(index.replace('<div class="chapter-grid">','<div class="chapter-grid">'+card),encoding='utf-8')
# Include unchanged dependencies only in the preview directory.
for rel in ['css/style.css','css/quiz-3d-securite.css']:
    shutil.copyfile(ROOT/rel,STAGE/rel)

def document(title):
    d=Document(); s=d.sections[0]
    for element in list(d.styles.element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.top_margin=Cm(1.5); s.bottom_margin=Cm(1.5); s.left_margin=Cm(1.7); s.right_margin=Cm(1.7)
    for sty in ['Normal','Title','Heading 1','Heading 2']:
        d.styles[sty].font.name='Arial'; d.styles[sty].font.color.rgb=RGBColor(0,0,0)
    normal=d.styles['Normal']; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(3)
    normal.paragraph_format.line_spacing=1.05
    d.styles['Title'].font.size=Pt(20)
    d.add_paragraph(title,'Title')
    foot=s.footer.paragraphs[0]; foot.alignment=2
    foot.add_run('3D SEGPA • Sciences physiques • ')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
    return d

d=document('Interrogation sur les transformations chimiques et la matière')
d.add_paragraph('Nom : ........................................   Prénom : ........................................')
d.add_paragraph('Classe : 3D SEGPA     Date : ........................     Note : ........ / 20')
p=d.add_paragraph(); p.add_run('Durée maximale : 20 minutes. ').bold=True
p.add_run('Coche une seule réponse par question. Chaque bonne réponse vaut 1 point. Une réponse fausse ou absente vaut 0 point.')
d.add_paragraph('Utilise les extraits du cours lorsqu’ils sont présents. Garde 2 minutes pour te relire.')
for i,q in enumerate(questions):
    if i in [5,10,15]: d.add_page_break()
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.keep_with_next=True
    p.add_run(f'{i+1}. {q["question"]}').bold=True
    if i in visuals:
        p=d.add_paragraph(); p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(images/visuals[i][0]),width=Cm(16.8))
    for j,a in enumerate(q['answers']):
        p=d.add_paragraph(f'☐ {chr(65+j)}. {a}')
        p.paragraph_format.left_indent=Cm(.3)
        p.paragraph_format.keep_with_next=j<2
d.save(OUT/'interrogation-3d-chimie.docx')

d=document('Corrigé de l’interrogation sur les transformations chimiques et la matière')
d.add_paragraph('3D SEGPA • 20 questions • Durée maximale : 20 minutes')
d.add_paragraph('Barème : 1 point par bonne réponse. Réponse fausse, absente ou multiple : 0 point. Total sur 20.')
for i,q in enumerate(questions):
    if i==10: d.add_page_break()
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.keep_with_next=True
    p.add_run(f'{i+1}. {chr(65+q["correct"])} — {q["answers"][q["correct"]]}').bold=True
    d.add_paragraph(q['explanation'])
d.save(OUT/'corrige-3d-chimie.docx')
print('Created 20-question website, 4 source crops, student DOCX and separate answer key.')
