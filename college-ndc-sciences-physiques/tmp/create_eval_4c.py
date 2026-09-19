from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path('output/documents'); out.mkdir(parents=True,exist_ok=True)
assets=Path('../images/qcm-4c')
d=Document(); sec=d.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5)
sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
    s=d.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    s.font.size=Pt(12 if name=='Normal' else 19 if name=='Title' else 14)
    s.paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.1
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Heading 1','Heading 2']:
    d.styles[name].paragraph_format.space_before=Pt(10)
    d.styles[name].paragraph_format.space_after=Pt(4)
d.core_properties.title='Évaluation sur les circuits électriques'
d.core_properties.subject='Physique chimie 4C SEGPA Rappels de 5e'
d.core_properties.author='Collège NDC'

def p(t='',bold=False):
    p=d.add_paragraph(); p.add_run(t).bold=bold; return p
def heading(t,points):
    d.add_heading(t,level=1); p('Barème : '+str(points)+' points',True)
def line():
    p('........................................................................................................')
def picture(par,name,width,alt):
    shape=par.add_run().add_picture(str(assets/(name+'.png')),width=Cm(width))
    shape._inline.docPr.set('descr',alt)
def table_headers(table):
    table.autofit=False
    for row in table.rows:
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margins.append(el)
            tcPr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            tcPr.append(borders)
    for c in table.rows[0].cells:
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'EDEDED'); c._tc.get_or_add_tcPr().append(sh)
        for r in c.paragraphs[0].runs:r.bold=True

d.add_paragraph('Évaluation sur les circuits électriques','Title')
p('Physique chimie • 4C SEGPA • Rappels de 5e')
p('Nom : ....................................  Prénom : ....................................')
p('Date : ....................................                         Note : ........ / 20')
p('Observe les images. Réponds avec des mots simples. Dessine au crayon et à la règle.')
heading('Exercice 1 Reconnaître le matériel',4)
t=d.add_table(rows=2,cols=2)
for i,(name,alt) in enumerate([('symbole-1','Un cercle contenant une croix.'),('symbole-2','Deux traits de longueurs différentes avec le signe plus.')]):
    t.cell(0,i).text='Symbole '+str(i+1)
    cell=t.cell(1,i); para=cell.paragraphs[0]; para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    picture(para,name,1.6,alt)
table_headers(t)
p('1. Écris le nom de chaque élément. (2 points)')
p('Symbole 1 : ................................   Symbole 2 : ................................')
p('2. Quel signe repère la borne positive de la pile ? (1 point)')
p('Réponse : ................................')
p('3. Sur le schéma A ci-dessous, combien y a-t-il de lampes ? (1 point)')
p('Réponse : ................................')
heading('Exercice 2 Lire les schémas',6)
t=d.add_table(rows=2,cols=2)
for i,(name,alt) in enumerate([('schema-a','Deux lampes placées sur le côté inférieur du tracé.'),('schema-b','Deux lampes sur deux traits horizontaux distincts.')]):
    t.cell(0,i).text='Schéma '+chr(65+i)
    para=t.cell(1,i).paragraphs[0]; para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    picture(para,name,5.3,alt)
table_headers(t)
p('1. Écris « en série » ou « en dérivation ». (4 points)')
p('Schéma A : ................................   Schéma B : ................................')
p('2. Quel schéma possède une seule boucle ? (1 point)   ...............')
p('3. Quel schéma a ses lampes sur deux branches différentes ? (1 point)   ...............')

d.add_page_break()
p('Nom : ....................................  Prénom : ....................................')
heading('Exercice 3 Prévoir ce qui se passe',6)
para=p(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
picture(para,'montage-a',7.6,'Montage du cours avec les lampes L2 à gauche et L1 à droite.')
p('1. On dévisse L1. Que devient L2 ? (2 points)')
line()
p('2. Dans le montage en dérivation du cours, on dévisse une lampe.')
p('Que devient l’autre lampe ? (2 points)')
line()
p('3. Complète avec « faible » ou « fort ». (2 points)')
p('Dans les montages du cours, les lampes ont :')
p('• un éclat ................................ dans le montage en série ;')
p('• un éclat ................................ dans le montage en dérivation.')
heading('Exercice 4 Dessiner un circuit',4)
p('Dessine un circuit en série avec une pile et deux lampes.')
p('Utilise les symboles du cours et ferme la boucle.')
p('Pile : 1 point • Deux lampes : 1 point • Montage en série : 1 point\nBoucle fermée : 1 point')
# Open drawing space, with no decorative box or table.
blank=p(); blank.paragraph_format.space_after=Pt(0)
blank.paragraph_format.line_spacing=1
blank.paragraph_format.space_before=Pt(110)
p('Relis tes réponses avant de rendre ta copie.')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('4C SEGPA • Circuits électriques • ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
d.save(out/'Evaluation_circuits_4C_SEGPA.docx')
print(out/'Evaluation_circuits_4C_SEGPA.docx')
